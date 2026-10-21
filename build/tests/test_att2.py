from tcommon import *
import openpyxl,shutil
async def main():
    async with async_playwright() as P:
        b=await P.chromium.launch();errs=[];res=[]
        def chk(n,c,extra=''):res.append(bool(c));print(('PASS ' if c else 'FAIL ')+n,extra if not c else '')
        ctx=await b.new_context(viewport={'width':1280,'height':900},accept_downloads=True);await ctx.add_init_script(MOCK);await ctx.add_init_script("window.showSaveFilePicker=undefined")
        pg=await ctx.new_page();pg.on('pageerror',lambda e:errs.append(str(e)))
        await pg.route('**/fonts.g*/**',lambda r:r.abort())
        async def cdn(r):
            if 'jszip' in r.request.url:await r.fulfill(body=JSZ,content_type='application/javascript')
            else:await r.abort()
        await pg.route('**/cdnjs**',cdn)
        await pg.goto('file:///home/claude/app/out/cloud.html');await pg.wait_for_timeout(700)
        await login(pg,'Jkgupta&Co','JKG@2500')
        for u in ['bhavika','kiran']:
            await pg.evaluate("staffF()");await pg.wait_for_timeout(300)
            await pg.fill('#au',u);await pg.fill('#ap','pass1234');await pg.click('.addf button:not(.hli)');await pg.wait_for_timeout(600)
        print(await pg.evaluate("Object.keys(U())"))
        await pg.evaluate("atTog(0)");await pg.evaluate("atTog(1)");await pg.wait_for_timeout(400)
        chk('both allowed attendance',await pg.evaluate("ATG().length==2"),await pg.evaluate("ATG()"))
        await pg.evaluate("logout()");await pg.wait_for_timeout(300)
        for u in ['bhavika','kiran']:
            await login(pg,u,'pass1234');chk(u+' logged in',await pg.evaluate("USER")==u)
            await pg.evaluate("att()");await pg.wait_for_timeout(400)
            await pg.evaluate("attIn()");await pg.wait_for_timeout(400)
            chk(u+' checked in',await pg.evaluate("!!(attGet(USER).d[ymd(new Date())]||{}).i"))
            await pg.evaluate("attOut()");await pg.wait_for_timeout(300)
            await pg.evaluate("logout()");await pg.wait_for_timeout(300)
        await login(pg,'Jkgupta&Co','JKG@2500');await pg.evaluate("att()");await pg.wait_for_timeout(800)
        txt=await pg.inner_text('#app');print(txt[:300])
        chk('admin sees both staff',('bhavika' in txt.lower()) and ('kiran' in txt.lower()))
        chk('admin data shows both present today',await pg.evaluate("['bhavika','kiran'].every(u=>(attGet(u).d[ymd(new Date())]||{}).s=='P')"))
        async with pg.expect_download(timeout=10000) as d: await pg.evaluate("attXls()")
        x=await d.value;shutil.copy(await x.path(),'out/att2.xlsx');ws=openpyxl.load_workbook('out/att2.xlsx')['Summary']
        names=[ws.cell(r,1).value for r in range(2,6)];print(names)
        chk('excel has both staff','bhavika' in names and 'kiran' in names)
        td=await pg.evaluate("+ymd(new Date()).slice(8)")
        chk("both marked P today in Excel",all(ws.cell(2+names.index(u),1+td).value=='P' for u in ['bhavika','kiran']))
        chk('no JS errors',not errs,errs)
        await b.close();print('ALL OK' if all(res) else 'SOME FAILED')
asyncio.run(main())
