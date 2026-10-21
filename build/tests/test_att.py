from tcommon import *
import openpyxl
async def main():
    async with async_playwright() as P:
        b=await P.chromium.launch();errs=[];res=[]
        def chk(n,c,extra=''):res.append(bool(c));print(('PASS ' if c else 'FAIL ')+n,extra if not c else '')
        ctx=await b.new_context(viewport={'width':1280,'height':900},accept_downloads=True);await ctx.add_init_script(MOCK)
        await ctx.add_init_script("window.showSaveFilePicker=undefined")
        pg=await ctx.new_page();pg.on('pageerror',lambda e:errs.append(str(e)))
        await pg.route('**/fonts.g*/**',lambda r:r.abort())
        async def cdn(r):
            if 'jszip' in r.request.url:await r.fulfill(body=JSZ,content_type='application/javascript')
            else:await r.abort()
        await pg.route('**/cdnjs**',cdn)
        await pg.goto('file:///home/claude/app/out/cloud.html');await pg.wait_for_timeout(700)
        await login(pg,'Jkgupta&Co','JKG@2500')
        await pg.evaluate("""(async()=>{const mo=attMonth();for(const u of ['ravi','sita']){await addUFake(u)}})""".replace("await addUFake(u)","0")) if False else None
        await pg.evaluate("""(()=>{const us=U();us.ravi='x';us.sita='x';LS.setItem('jkgusers',JSON.stringify(us))})()""")
        print(await pg.evaluate("Object.keys(U())"))
        mo=await pg.evaluate("attMonth()")
        await pg.evaluate("""(mo=>{for(const [u,days] of [['ravi',[['01','P','09:30','18:00'],['02','P','10:00','17:30'],['03','A','',''],['04','H','10:00','14:00'],['05','L','','']]],['sita',[['01','P','09:00','18:00']]]]){const o=attGet(u);days.forEach(([d,s,i,t])=>{o.d[mo+'-'+d]={s,i,o:t}});attSave(u,o)}})""",)  if False else None
        await pg.evaluate("""mo=>{window.__mo=mo}""",mo)
        await pg.evaluate("""()=>{const mo=attMonth();const set=(u,d,s,i,t)=>{const o=attGet(u);o.d[mo+'-'+d]={s,i,o:t};attSave(u,o)};
          set('ravi','01','P','09:30','18:00');set('ravi','02','P','10:00','17:30');set('ravi','03','A','','');set('ravi','04','H','10:00','14:00');set('ravi','05','L','','');set('sita','01','P','09:00','18:00')}""")
        await pg.evaluate("att()");await pg.wait_for_timeout(500)
        async with pg.expect_download(timeout=10000) as d: await pg.evaluate("attXls()")
        x=await d.value;fn=x.suggested_filename;p=await x.path();print(fn)
        chk('admin export is .xlsx',fn.endswith('.xlsx'))
        import shutil;shutil.copy(p,'/home/claude/app/out/att_admin.xlsx')
        import shutil;shutil.copy(p,'out/att_admin.xlsx');wb=openpyxl.load_workbook('out/att_admin.xlsx');print(wb.sheetnames)
        ws=wb['Summary'];print([c.value for c in ws[1]][:8],[c.value for c in ws[2]][:8])
        chk('two sheets',wb.sheetnames==['Summary','Daily details'])
        chk('staff rows',{ws['A2'].value,ws['A3'].value}=={'ravi','sita'},[ws['A2'].value,ws['A3'].value])
        r=[r for r in range(2,4) if ws.cell(r,1).value=='ravi'][0]
        chk('status letters kept',[ws.cell(r,c).value for c in range(2,7)]==['P','P','A','H','L'],[ws.cell(r,c).value for c in range(2,7)])
        nc=ws.max_column
        chk('totals are formulas',str(ws.cell(r,nc-5).value).startswith('=COUNTIF'),ws.cell(r,nc-5).value)
        chk('has conditional formatting',len(ws.conditional_formatting)>0)
        chk('has dropdown validation',len(ws.data_validations.dataValidation)>0)
        chk('frozen panes',ws.freeze_panes=='B2',ws.freeze_panes)
        chk('header colour',ws['A1'].fill.fgColor.rgb=='FF12305A')
        dd=wb['Daily details'];chk('daily rows',dd.max_row>=60,dd.max_row);print([c.value for c in dd[2]])
        # staff self export
        await pg.evaluate("tg(1)") ;await pg.evaluate("logout()");await pg.wait_for_timeout(300)
        await login(pg,'ravi','JKG1234');await pg.evaluate("att()");await pg.wait_for_timeout(500)
        async with pg.expect_download(timeout=10000) as d: await pg.evaluate("attXlsSelf()")
        x=await d.value;p=await x.path();shutil.copy(p,'out/att_self.xlsx');wb=openpyxl.load_workbook('out/att_self.xlsx');ws=wb.worksheets[0];print(x.suggested_filename,ws['A1'].value)
        chk('self export title',ws['A1'].value and 'ravi' in ws['A1'].value)
        chk('self export formula',any(str(c.value).startswith('=COUNTIF') for row in ws.iter_rows() for c in row))
        chk('no JS errors',not errs,errs)
        await b.close();print('ALL OK' if all(res) else 'SOME FAILED')
asyncio.run(main())
